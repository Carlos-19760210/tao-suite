# -*- coding: utf-8 -*-
# Parecer Técnico da RT — checklist de validação do TAO Lab para o corte FCerta->TAO (DOCX).
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT   = r"C:\Users\carlo\Parecer_RT_TAO_Lab.docx"
AZUL  = RGBColor(0x1e, 0x40, 0xaf)
CINZA = RGBColor(0x64, 0x74, 0x8b)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.8); s.left_margin = s.right_margin = Cm(1.8)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)

def h1(t):
    p = doc.add_heading(t, level=1); p.runs[0].font.color.rgb = AZUL; return p
def h2(t):
    p = doc.add_heading(t, level=2); p.runs[0].font.color.rgb = AZUL; return p
def para(t, italic=False, size=11):
    p = doc.add_paragraph(); r = p.add_run(t); r.italic = italic; r.font.size = Pt(size); return p

# ── CAPA ──
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('\n\nParecer Técnico da\nFarmacêutica Responsável Técnica'); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = AZUL
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('Validação do sistema TAO Lab para substituição do Formula Certa'); r.font.size = Pt(14)
s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s2.add_run('Conformidade RDC 67/2007 · Portaria 344/98 · SNGPC'); r.font.size = Pt(12); r.font.color.rgb = CINZA

doc.add_paragraph('\n')
para('Farmácia: Magis-TAO Farmácia de Manipulação — Cotia/SP', size=11)
para('Data da avaliação: ____ / ____ / ________', size=11)
para('Este documento reúne os pontos do sistema que dependem da validação técnica da RT antes do '
     'corte definitivo (desligamento do Formula Certa). Para cada item, assinale "De acordo" ou '
     '"Ajustar" e registre observações quando necessário.', italic=True, size=10)
doc.add_page_break()

# (área, [ (titulo, descricao) ... ])
AREAS = [
 ("A. Cálculos e Motor Farmacotécnico", [
   ("Correção dos cálculos", "Fator de correção, perda, diluição, teor, densidade, equivalência e dose. Teste de ouro de 99,92% contra o Formula Certa; a RT valida a farmacotécnica."),
   ("Fator de equivalência sal↔base", "Recalculado pelos pesos moleculares (ex.: nortriptilina 1,14; escitalopram 1,28; terbinafina 1,12). A RT confirma os fatores."),
   ("Alerta de dose máxima (não-bloqueante)", "O sistema alerta em vermelho quando a dose ultrapassa o máximo, mas NÃO impede (a posologia pode dividir a dose). A RT confirma o comportamento."),
   ("Trava de restrição", "Substâncias bloqueadas/restritas (ex.: semaglutida) são recusadas no orçamento. A RT valida a lista."),
 ]),
 ("B. Prazos de Validade (padrões adotados)", [
   ("Validade das preparações", "Florais 90 dias; demais preparações 120 dias."),
   ("Validade das diluições/bases produzidas internamente", "90 dias a partir da produção."),
 ]),
 ("C. Rótulo e Livro de Receituário", [
   ("Rótulo (RDC 67, Anexo I)", "A RT confere se todos os dizeres obrigatórios estão presentes e corretos (paciente, prescritor, composição, lote/validade, posologia, farmácia+RT, advertências, conservação)."),
   ("Livro de Receituário", "Registro sequencial das manipulações — formato atende à Lei 5.991 art. 42 e à RDC 67."),
 ]),
 ("D. Controlados / SNGPC (Portaria 344/98)", [
   ("Classificação dos controlados", "39 ativos marcados como controlados com a classe (A1/A2/B1/B2/C1/C2/C5/AM) importada do Formula Certa. A RT valida a classe de cada substância."),
   ("Escrituração da transformação de controlados", "Ao produzir internamente um diluído controlado (ex.: clonazepam 1:50), o sistema hoje NÃO gera movimento SNGPC (evita dupla contagem; a saída é escriturada na dispensação da OM). A RT confirma se atende à fiscalização, ou se a transformação deve ser escriturada."),
   ("Notificação de receita e identificação do comprador", "Campos exigidos na OM controlada (tipo de receita, nº de notificação, comprador). A RT valida."),
 ]),
 ("E. Qualidade e Especificações", [
   ("Controle de qualidade de recebimento", "Critérios para aprovar/reprovar lote na entrada (quem executa, o que confere)."),
   ("Ficha técnica / especificação da matéria-prima", "Especificações preenchidas (fórmula/peso molecular via fonte pública + dados manuais). A RT valida."),
   ("Peso médio de cápsulas", "Mantido FORA do sistema (inspeção manual da RT). A RT confirma a decisão."),
   ("Fórmulas padrão", "734 fórmulas padrão importadas do Formula Certa. A RT valida as que permanecem em uso."),
 ]),
 ("F. Comportamento Operacional", [
   ("Baixa de estoque na conclusão da OM", "O estoque dos lotes pesados é baixado uma única vez, na etapa final da produção. Modelo aprovado?"),
   ("Alerta de licença de fornecedor vencida (não-bloqueante)", "No recebimento de NF, se o fornecedor tiver AFE/AE/licença vencida, o sistema avisa mas não bloqueia. Aceitável?"),
   ("Etapas do kanban de produção", "As etapas padrão refletem o fluxo real da bancada?"),
 ]),
 ("G. Escopo de Manipulação (a RT formaliza)", [
   ("Substâncias não manipuladas", "A Magis NÃO manipula: GLP-1, estéreis, homeopatia, antibióticos (antimicrobianos) e não opera Farmácia Popular. A RT confirma o escopo."),
 ]),
]

n = 0
for area, itens in AREAS:
    h2(area)
    tb = doc.add_table(rows=1, cols=4); tb.style = 'Light Grid Accent 1'
    hd = tb.rows[0].cells
    for i, txt in enumerate(['#', 'Ponto a validar', 'De acordo', 'Ajustar / Observações']):
        hd[i].paragraphs[0].add_run(txt).bold = True
    for titulo, desc in itens:
        n += 1
        c = tb.add_row().cells
        c[0].text = str(n)
        c[1].paragraphs[0].add_run(titulo).bold = True
        c[1].add_paragraph(desc).runs[0].font.size = Pt(9)
        c[2].text = '☐ Sim'
        c[3].text = '☐ Ajustar: ______________________'
    for row in tb.rows:
        row.cells[0].width = Cm(0.9); row.cells[1].width = Cm(9.2); row.cells[2].width = Cm(2.2); row.cells[3].width = Cm(5.0)
    doc.add_paragraph()

# ── PARECER FINAL + ASSINATURA ──
doc.add_page_break()
h1('Parecer final')
para('☐  APROVADO — o sistema atende aos requisitos técnicos e legais avaliados; autorizo o uso em produção.')
para('☐  APROVADO COM RESSALVAS — atende, condicionado aos ajustes assinalados acima.')
para('☐  NÃO APROVADO — pendências impedem o uso; ver observações.')
doc.add_paragraph('\n')
para('Observações gerais:')
for _ in range(4):
    doc.add_paragraph('_' * 95)
doc.add_paragraph('\n\n')
para('__________________________________________', size=11)
para('Farmacêutica Responsável Técnica', size=10)
para('Nome: ______________________________________', size=11)
para('CRF: __________ / UF: ______     Data: ____ / ____ / ________', size=11)

doc.save(OUT)
print('DOCX gerado:', OUT, '-', os.path.getsize(OUT)//1024, 'KB', '-', n, 'pontos')
