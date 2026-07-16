# -*- coding: utf-8 -*-
# Manual do Usuário do TAO CRM (produto vendável) — DETALHADO, tela a tela.
# Reaproveita "O card do cliente" do módulo compartilhado manual_cap_card.py (incluir_formula=False).
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from manual_cap_card import capitulo_card

PRINTS = r"C:\Users\carlo\AppData\Local\Temp\claude\C--Users-carlo\fc46e316-08dd-4693-9e35-805a6cb6dffe\scratchpad\prints"
OUT    = r"C:\Users\carlo\Manual_Usuario_TAO_CRM_v1.docx"
AZUL   = RGBColor(0x1e, 0x40, 0xaf)
CINZA  = RGBColor(0x64, 0x74, 0x8b)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.8); s.left_margin = s.right_margin = Cm(1.8)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)

def h1(t):
    p = doc.add_heading(t, level=1); p.runs[0].font.color.rgb = AZUL; return p
def h2(t):
    p = doc.add_heading(t, level=2); p.runs[0].font.color.rgb = AZUL; return p
def h3(t):
    p = doc.add_heading(t, level=3); p.runs[0].font.color.rgb = AZUL; return p
def para(t, italic=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.italic = italic; return p
def passo(t):
    doc.add_paragraph(t, style='List Number')
def nota(t):
    p = doc.add_paragraph(); r = p.add_run('Regra: '); r.bold = True; r.font.color.rgb = AZUL; p.add_run(t)
def imagem(nome, legenda=None, larg=16.8):
    fp = os.path.join(PRINTS, nome + '.png')
    if os.path.isfile(fp):
        doc.add_picture(fp, width=Cm(larg))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if legenda:
            c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run('Figura — ' + legenda); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = CINZA
def tabela_campos(linhas, col0='Item'):
    t = doc.add_table(rows=1, cols=3); t.style = 'Light Grid Accent 1'
    hdr = t.rows[0].cells
    for i, txt in enumerate([col0, 'Obrig.', 'Descrição e regra']):
        hdr[i].paragraphs[0].add_run(txt).bold = True
    for campo, ob, desc in linhas:
        c = t.add_row().cells
        c[0].paragraphs[0].add_run(campo).bold = True
        c[1].text = 'Sim' if ob else '—'
        c[2].text = desc
    for row in t.rows:
        row.cells[0].width = Cm(3.8); row.cells[1].width = Cm(1.6); row.cells[2].width = Cm(11.4)
    doc.add_paragraph()

_H = { 'h1': h1, 'h2': h2, 'h3': h3, 'para': para, 'passo': passo, 'nota': nota,
       'tabela_campos': lambda l: tabela_campos(l), 'imagem': imagem, 'page_break': lambda: doc.add_page_break() }

# ── CAPA ──
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('\n\n\nManual do Usuário'); r.bold = True; r.font.size = Pt(32); r.font.color.rgb = AZUL
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('TAO CRM — Atendimento e Vendas por WhatsApp'); r.font.size = Pt(18)
s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s2.add_run('Kanban · Chat WhatsApp · Automações · Campanhas · Indicadores'); r.font.size = Pt(13); r.font.color.rgb = CINZA
doc.add_page_break()

# ── APRESENTAÇÃO ──
h1('Apresentação')
para('O TAO CRM organiza todo o atendimento e as vendas da sua empresa em um funil visual (Kanban), com o chat de WhatsApp integrado. Cada cliente vira um card, e é nele que a equipe conversa, registra o combinado, acompanha a negociação e fecha a venda — sem trocar de tela.')
para('Este manual descreve, tela a tela e campo a campo, como operar o CRM: o Kanban, o card do cliente, os contatos, as configurações do funil (pipelines, estágios e campos), as automações, as mensagens e os indicadores.')
h3('Convenções')
para('• Campos marcados com * são obrigatórios.')
para('• Onde houver "Regra:", trata-se de um comportamento automático ou validação do sistema.')
para('• O contato (cliente) é ÚNICO em toda a solução — o mesmo do card, do Agente de WhatsApp e das Campanhas.', italic=True)

# ── PRIMEIROS PASSOS ──
doc.add_page_break(); h1('Primeiros passos — acesso e navegação')
para('O TAO CRM é usado pelo navegador (Chrome, Edge ou Firefox), no computador ou no celular — não há nada para instalar.')
h3('Como entrar')
passo('Abra o endereço do portal (fornecido pelo administrador) e faça login com e-mail e senha.')
passo('A tela inicial mostra o menu CRM (Painel e Kanban) e, para gestores, o menu Configurações.')
nota('O acesso é por perfil: o gestor vê todos os cards e as Configurações; o atendente vê apenas os cards sob sua responsabilidade.')
h3('Como navegar')
para('O menu CRM tem dois itens de operação — Painel (indicadores) e Kanban (o funil) — e, para gestores, Configurações (funil, campos, automações, equipe, integração).')
nota('Em qualquer lista há uma busca no topo; nos campos que buscam pessoas, navegue pelo teclado (↑↓/Enter).')

# ── VISÃO GERAL ──
doc.add_page_break(); h1('Como funciona — o funil de atendimento')
para('Entenda o caminho de um atendimento antes de entrar nas telas. Cada etapa é detalhada adiante.')
passo('CHEGA O CONTATO. O cliente manda mensagem no WhatsApp; o sistema cria (ou reabre) um card no funil, já com o contato identificado pelo número.')
passo('ATENDIMENTO. A equipe conversa pelo card, registra notas internas e vai movendo o card pelas fases (ex.: Aguardando Atendimento → Em Conversa → Orçamento Enviado → Negociação).')
passo('FECHAMENTO. O card é fechado como GANHO (venda) ou PERDIDO (com motivo). Ganhos podem cruzar automaticamente para um funil de Pós-vendas.')
passo('PÓS E RETENÇÃO. Automações e campanhas cuidam do follow-up, do NPS e da reativação de clientes.')
nota('Em uma linha: Contato → Atendimento (card + WhatsApp) → Ganho/Perdido → Pós-venda. Tudo acontece no card.')

# ── O CARD (capítulo reutilizado — SEM a parte de fórmula) ──
capitulo_card(_H, incluir_formula=False)

# ── KANBAN ──
doc.add_page_break(); h1('O Kanban')
para('Menu CRM → Kanban. A visão em colunas de todo o funil: cada coluna é uma fase; cada cartão, um atendimento. É a tela onde a equipe trabalha o dia inteiro.')
imagem('crm_kanban', 'O Kanban — funil em colunas, filtros e busca (dados dos clientes ocultados)')
h3('Trabalhar os cards')
passo('Clique em um card para abrir a ficha; ARRASTE-o entre as colunas para mudar a fase.')
passo('Ao mover para uma coluna de Ganho/Perdido (ou pelos botões da ficha), o sistema pede a confirmação do valor (ganho) ou o motivo (perdido).')
passo('Use "+ Novo Card" para abrir um atendimento manualmente; "Inbox" mostra as conversas com mensagens não lidas.')
h3('Buscar e filtrar')
passo('A busca do topo localiza por nome, WhatsApp ou número da requisição.')
passo('Os filtros (Atendente, Fase, Status, Tag) focam a visão; "Mostrar colunas encerradas" exibe os ganhos/perdidos.')
passo('O intervalo de atualização automática é configurável (10 s a desligado). Alterne entre os funis pelas abas (Funil de Vendas / Pós-vendas).')
nota('Cada alteração num card define quem o alterou como responsável. Um ponto vermelho no card indica mensagem não lida.')

# ── CONTATOS ──
doc.add_page_break(); h1('Contatos')
para('Menu Cadastros → Clientes/Contatos. A base ÚNICA de pessoas — o mesmo contato do card, do Agente e das Campanhas. Não crie cadastros paralelos.')
imagem('crm_contatos', 'Lista de contatos (WhatsApp ocultado)')
passo('Busque por nome, WhatsApp ou e-mail; filtre por workspace; "+ Novo Contato" para cadastrar.')
tabela_campos([
 ('Nome', True, 'Nome do contato.'),
 ('WhatsApp', True, 'CHAVE ÚNICA do cadastro — o sistema não duplica pessoas com o mesmo número.'),
 ('E-mail / Cidade', False, 'Dados de contato e localização.'),
 ('Classificação', False, 'Marcadores do relacionamento (ex.: cliente, lead).'),
 ('Endereço completo', False, 'Usado por módulos como Entregas.'),
])
nota('Ao salvar, os dados valem para o CRM, o Agente e as Campanhas — é o mesmo contato.')

# ── PIPELINES E ESTÁGIOS ──
doc.add_page_break(); h1('Pipelines e Estágios')
para('Configurações → Pipelines. Você define os funis (pipelines) e as fases (estágios) de cada um. É a espinha dorsal do CRM.')
imagem('crm_pipelines', 'Configuração de pipelines e estágios')
passo('Crie um pipeline pelo campo "Novo pipeline" (nome + ordem + Criar). O 1º pipeline costuma ser o Funil de Vendas; o 2º, o Pós-vendas.')
passo('No pipeline selecionado, adicione/renomeie/reordene os estágios e escolha a COR de cada um.')
passo('Marque "Usar pipeline selecionado como Pós-vendas" para que os cards GANHOS cruzem automaticamente para ele.')
h3('Tipo do estágio')
para('O TIPO de cada estágio controla o comportamento do card ao chegar nele:')
tabela_campos([
 ('Normal', False, 'Fase comum de andamento (ex.: Em Conversa, Negociação).'),
 ('Ganho', False, 'Fecha o card como venda. Ao entrar, dispara os módulos ligados (venda, entrega, etc.).'),
 ('Perdido', False, 'Fecha o card como perdido — exige motivo.'),
 ('Handoff', False, 'Fase de entrada do atendimento humano (onde o card chega quando o cliente pede uma pessoa).'),
])
nota('Cards ganhos no Funil de Vendas cruzam para a 1ª fase do pipeline marcado como Pós-vendas. Os cards apontam para o ID da fase — renomear/reordenar não perde cards.')

# ── CAMPOS POR FASE ──
doc.add_page_break(); h1('Campos personalizados (checklist por fase)')
para('Configurações → Campos. Crie campos próprios e associe a cada fase — é como você monta o checklist de conferência que o card exige.')
imagem('crm_campos', 'Campos personalizados e associação por estágio')
passo('Crie um campo definindo o Nome, o Tipo (texto, número, Sim/Não, lista de opções) e, se for lista, as opções.')
passo('Associe o campo a um ou mais estágios. Para cada associação, marque "na entrada" (aparece ao chegar na fase) e "obrigatório" (exige preenchimento).')
nota('É esse checklist que aparece no card (faixa âmbar): o card não avança de fase — nem fecha como ganho — enquanto houver campo obrigatório vazio na fase atual.')

# ── AUTOMAÇÕES ──
doc.add_page_break(); h1('Automações')
para('Configurações → Automações. Regras que agem sozinhas sobre os cards, sem intervenção da equipe.')
imagem('crm_automacoes', 'Cadastro de automações (gatilho → condição → ação)')
para('Cada regra é: um GATILHO + uma condição (fase/tempo) + uma AÇÃO. Ative para valer.')
h3('Gatilhos')
tabela_campos([
 ('Entrar na fase', False, 'Dispara quando o card chega a um estágio.'),
 ('Tempo na fase', False, 'Dispara após X minutos/horas parado na fase (ex.: 72 h).'),
 ('Recebeu mensagem', False, 'Dispara quando o cliente responde.'),
 ('Enviou mensagem', False, 'Dispara quando o atendente responde pelo CRM.'),
 ('Sem resposta', False, 'Dispara após um período sem resposta do cliente.'),
])
h3('Ações')
tabela_campos([
 ('Enviar mensagem', False, 'Manda uma mensagem de WhatsApp (usa um template/texto).'),
 ('Mover de fase', False, 'Move o card para outro estágio.'),
 ('Atribuir responsável', False, 'Define/rodízio (round-robin) do responsável.'),
 ('Notificar e-mail', False, 'Avisa a equipe por e-mail.'),
 ('Fechar como perdido', False, 'Fecha o card com um motivo padrão.'),
])
nota('Exemplo clássico: gatilho "Tempo na fase = 72 h" na coluna "Aguarda Resposta" → ação "Fechar como perdido" com o motivo "Não responde os contatos".')

# ── ORGANIZAÇÃO (tags, lembretes, comentários) ──
doc.add_page_break(); h1('Tags, Lembretes e Comentários')
para('Recursos do card para organizar o atendimento e não perder o timing (também citados no capítulo do card). As Tags são cadastradas em Configurações → Tags.')
tabela_campos([
 ('Tags', False, 'Etiquetas coloridas por workspace, para classificar e filtrar cards no Kanban.'),
 ('Lembretes', False, 'Follow-up com data/hora e notificação por e-mail quando vence.'),
 ('Comentários internos', False, 'Notas da equipe no card; o cliente não vê.'),
])

# ── MENSAGENS E TEMPLATES ──
doc.add_page_break(); h1('Mensagens, Templates e Agendamento')
para('A conversa do card usa o WhatsApp integrado (via a instância configurada na Integração). Para agilizar, use templates e agendamento.')
imagem('crm_templates', 'Templates de mensagem')
passo('Templates (Configurações → Templates): cadastre mensagens prontas para reutilizar no atendimento e nas automações.')
passo('Agendamento: no card, programe uma mensagem para uma data/hora futura (o sistema envia sozinho).')
nota('O envio respeita o horário de atendimento configurado (Configurações → Horário).')

# ── PAINEL ──
doc.add_page_break(); h1('Painel e Indicadores')
para('Menu CRM → Painel. Os números do funil, com filtro de período (hoje, 7, 30, 90 dias).')
imagem('crm_painel', 'Painel de indicadores do CRM')
tabela_campos([
 ('Cards abertos / Novos leads', False, 'Volume no pipeline e entradas no período.'),
 ('Taxa de conversão', False, 'Ganhos ÷ decididos no período.'),
 ('Em aberto / Receita gerada', False, 'Valor das oportunidades ativas e receita dos ganhos.'),
 ('Tempo até ganho / TMR', False, 'Tempo médio criação→ganho e tempo da 1ª resposta.'),
 ('Perdas por motivo', False, 'Ranking dos motivos de perda (qtde e valor).'),
 ('NPS / Renovações', False, 'Satisfação e eficiência da retenção (quando há dados no período).'),
])
nota('O painel também mostra o status das conexões de WhatsApp e o total de cards aguardando atendimento humano (handoff).')

# ── EQUIPE E PERMISSÕES ──
doc.add_page_break(); h1('Equipe e Permissões')
para('Configurações → Equipe. Define quem é gestor e quem é atendente/vendedor no workspace.')
imagem('crm_equipe', 'Equipe do workspace')
nota('Gestor vê todos os cards do workspace e acessa as Configurações; o atendente vê apenas os cards sob sua responsabilidade. Os selects de responsável mostram só a equipe do negócio.')

# ── CONFIGURAÇÕES / INTEGRAÇÃO ──
doc.add_page_break(); h1('Configurações e Integração')
para('Configurações → Integração. Conexão do WhatsApp (instâncias), webhooks de saída, e a chave de recebimento de mensagens.')
imagem('crm_integracao', 'Integração — instâncias de WhatsApp e webhooks')
passo('Conecte a(s) instância(s) de WhatsApp (Evolution) — o Kanban recebe as mensagens por elas.')
passo('Configure os webhooks de saída (por evento) quando quiser integrar com sistemas externos.')
passo('Defina o horário de atendimento (Configurações → Horário) — automações e envios o respeitam.')
nota('Outras abas de Configurações: Workspaces (negócios), Planos, SLA, CSAT/NPS, Metas, LGPD e Logs.')

doc.save(OUT)
print('DOCX gerado:', OUT, '-', os.path.getsize(OUT)//1024, 'KB')
