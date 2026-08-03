# -*- coding: utf-8 -*-
# Protocolo de Validação Farmacotécnica do sistema (RDC 67) — DOCX para a RT revisar,
# executar os casos de teste e ASSINAR. Modelo; o de-acordo técnico é da RT.
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT  = os.environ.get("PROT_OUT", r"C:\Users\carlo\Protocolo_Validacao_Farmacotecnica_v1.docx")
AZUL = RGBColor(0x1e,0x40,0xaf); CINZA = RGBColor(0x64,0x74,0x8b)

doc = Document()
for s in doc.sections:
    s.top_margin=s.bottom_margin=Cm(1.8); s.left_margin=s.right_margin=Cm(1.8)
doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(11)

def h1(t): p=doc.add_heading(t,level=1); p.runs[0].font.color.rgb=AZUL
def h2(t): p=doc.add_heading(t,level=2); p.runs[0].font.color.rgb=AZUL
def P(t="",b=False): p=doc.add_paragraph(); r=p.add_run(t); r.bold=b; return p
def campo(rot):
    p=doc.add_paragraph(); r=p.add_run(rot+': '); r.bold=True; p.add_run('_'*40)

# ── CAPA ──
t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=t.add_run('\n\nProtocolo de Validação Farmacotécnica'); r.bold=True; r.font.size=Pt(26); r.font.color.rgb=AZUL
s=doc.add_paragraph(); s.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=s.add_run('Sistema de manipulação — conformidade RDC 67/2007'); r.font.size=Pt(14); r.font.color.rgb=CINZA
s2=doc.add_paragraph(); s2.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=s2.add_run('Versão do protocolo v1 — a ser executado e assinado pela Responsável Técnica'); r.italic=True; r.font.size=Pt(11)
doc.add_paragraph()

h1('1. Identificação')
for c in ['Farmácia / razão social','CNPJ','Responsável Técnica (RT)','CRF','Sistema avaliado / versão','Data da validação']:
    campo(c)

h1('2. Objetivo')
P('Atestar, mediante execução de casos de teste, que o sistema realiza corretamente os cálculos '
  'farmacotécnicos e mantém os controles exigidos pela RDC 67/2007 e pela Portaria 344/98, de modo que '
  'possa ser utilizado como ferramenta de apoio à manipulação, SEMPRE sob revisão da RT.')

h1('3. Escopo')
P('São validadas as funções críticas abaixo. Para cada caso, a RT executa a operação no sistema, '
  'compara o Resultado obtido com o Resultado esperado e registra a conformidade (C = conforme / NC = não conforme), '
  'com observações quando necessário.')

CASOS = [
 ('Cálculo da quantidade a pesar', 'Fórmula com dose e nº de unidades definidos (ex.: Vitamina C 500 mg, 60 cápsulas).',
  'Qtd a pesar = dose × unidades ÷ (teor/100) × equivalência × diluição. Ex.: 500 mg × 60 = 30 g.'),
 ('Fator de correção pelo teor do lote', 'Insumo com teor do lote < 100% (ex.: 95%).',
  'A quantidade a pesar aumenta para compensar (ex.: 6 g ÷ 0,95 = 6,32 g), quando a chave de recálculo pelo lote está ativa.'),
 ('Equivalência sal ↔ base', 'Ativo prescrito como base e insumo comprado como sal (fator de equivalência cadastrado).',
  'A quantidade a pesar é ajustada pelo fator (ex.: × 1,15).'),
 ('Diluição (1:N)', 'Insumo diluído de fábrica (ex.: 1:10).',
  'A quantidade a pesar é multiplicada pelo fator de diluição.'),
 ('Escolha do lote (FEFO / em uso)', 'Ativo com mais de um lote liberado em estoque.',
  'O sistema seleciona automaticamente o lote em uso (aberto) ou, na falta, o de validade mais próxima; lote bloqueado/quarentena não é oferecido.'),
 ('Validade da preparação', 'OM concluída usando um lote cuja validade é menor que o prazo da forma.',
  'A validade da OM/rótulo é a MENOR entre o prazo da forma e a validade do lote, com alerta.'),
 ('Controle de qualidade do lote', 'Lote recém-recebido por NF.',
  'Nasce em quarentena; só entra na produção após aprovação da RT (não é possível usar lote não aprovado).'),
 ('Laudo por lote (Certificado de Análise)', 'Importação dos laudos de uma NF.',
  'O casamento laudo → lote é apresentado para conferência e só é gravado após confirmação da RT; o PDF fica arquivado.'),
 ('Gate de aprovação do orçamento', 'Orçamento novo em um card.',
  'Nasce "pendente de revisão"; aprovar é exclusivo do farmacêutico; o negócio não fecha como ganho sem ao menos um orçamento aprovado.'),
 ('Escrituração de controlados (SNGPC)', 'OM com componente controlado (Portaria 344/98).',
  'A OM exige prescritor, comprador e notificação; ao concluir, gera a saída no livro de controlados automaticamente.'),
 ('Baixa de estoque', 'Conclusão de uma OM.',
  'O estoque dos lotes pesados é baixado uma única vez, no momento da conclusão (não na geração nem na pesagem).'),
 ('Rótulo (RDC 67)', 'OM pronta.',
  'O rótulo traz os dados exigidos (composição, lote, validade, RT/CRF, cuidados) e a validade correta.'),
]
tbl=doc.add_table(rows=1, cols=5); tbl.style='Light Grid Accent 1'
hd=tbl.rows[0].cells
for i,x in enumerate(['#','Função','Cenário','Resultado esperado','C / NC']):
    hd[i].paragraphs[0].add_run(x).bold=True
for i,(fn,cen,esp) in enumerate(CASOS,1):
    c=tbl.add_row().cells
    c[0].text=str(i); c[1].paragraphs[0].add_run(fn).bold=True; c[2].text=cen; c[3].text=esp; c[4].text=''
for row in tbl.rows:
    row.cells[0].width=Cm(0.8); row.cells[1].width=Cm(3.4); row.cells[2].width=Cm(4.6); row.cells[3].width=Cm(5.6); row.cells[4].width=Cm(1.6)

doc.add_paragraph()
h1('4. Observações da RT')
for _ in range(4): P('_'*95)

h1('5. Conclusão e de-acordo')
P('Declaro que executei os casos de teste acima e que o sistema, na versão avaliada, atende aos requisitos '
  'farmacotécnicos verificados, para uso como ferramenta de apoio à manipulação sob minha supervisão. '
  'As não conformidades eventuais estão registradas no item 4 e devem ser tratadas antes do uso em produção.')
doc.add_paragraph(); doc.add_paragraph()
campo('Resultado geral (Aprovado / Aprovado com ressalvas / Reprovado)')
doc.add_paragraph()
P('___________________________________________')
P('Responsável Técnica — nome e assinatura')
campo('CRF'); campo('Data')

doc.save(OUT)
print('Protocolo gerado:', OUT, '-', os.path.getsize(OUT)//1024, 'KB')
