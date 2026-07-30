#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CORREÇÃO DIRECIONADA — pesagem de ativos DILUÍDOS ("1:N" / Fórmula Padrão) em OMs APROVADAS.
Corrige os dois defeitos herdados de imports antigos:
  (A) diluição não aplicada (ex.: VIT D3 1:100 em UI pesou ÷100);
  (B) variante ERRADA associada (ex.: PICOLINATO 1:100 casou o cadastro dil=1 em vez do dil=100).

REDE DE SEGURANÇA: recalcula qtd_total_g pela regra do motor e SÓ marca p/ gravar o item cujo
recálculo BATE com o FCerta (QTREAL de FC12110). Divergente do FCerta = revisão manual (não grava).

DRY-RUN por padrão (só gera planilha). --commit grava:
  • orcamentos.itens[i] (qtd_total_g, subtotal, ativo_id/codigo_fc/diluicao/teor/... se re-casou)
  • lab_ordem_itens.qtd_pesar da OM correspondente
Uso: python corrige_diluidos.py --db "C:\\...\\fcerta_29.ib" [--commit]
"""
import fdb, json, urllib.request, urllib.error, argparse, re, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

FB_DLL = r"C:\Users\carlo\FCertaSync\fb25\fbembed.dll"
SB  = "https://gclayesytzzpzkjvgede.supabase.co/rest/v1"
KEY = "sb_secret_HpoqM6ujk2yD6la7KM3cuQ_pdWBK8jo"
CID = "62f98634-77ff-42f4-acaf-8561d56583da"
APROVADOS = ("aprovado_farma", "aceito_paciente")

def dec(v): return v.decode("latin-1","replace").strip() if isinstance(v,(bytes,bytearray)) else (v.strip() if isinstance(v,str) else v)
def f(v, d=0.0):
    try: return float(v)
    except (TypeError, ValueError): return d
def serier_int(v):
    if isinstance(v,int): return v
    t=(dec(v) or "0").upper()
    if t.isdigit(): return int(t)
    if len(t)==1 and "A"<=t<="Z": return 10+ord(t)-ord("A")
    return 0
def sb_all(path):
    out,page=[],0
    while True:
        h={"apikey":KEY,"Authorization":"Bearer "+KEY,"Range-Unit":"items","Range":f"{page*1000}-{page*1000+999}"}
        ch=json.loads(urllib.request.urlopen(urllib.request.Request(SB+path,headers=h)).read().decode())
        out+=ch
        if len(ch)<1000: return out
        page+=1
def sbq(p):
    h={"apikey":KEY,"Authorization":"Bearer "+KEY}
    return json.loads(urllib.request.urlopen(urllib.request.Request(SB+p,headers=h)).read().decode())
def sb_patch(path, body):
    r=urllib.request.Request(SB+path, data=json.dumps(body).encode(), method="PATCH",
        headers={"apikey":KEY,"Authorization":"Bearer "+KEY,"Content-Type":"application/json","Prefer":"return=minimal"})
    return urllib.request.urlopen(r).status

DIL_RE = re.compile(r"(?<!\d)1:(\d{1,4})(?!\d)")
def toks(s): return [t for t in re.split(r"\s+", (s or "").strip()) if t]

def recalc_qtd_g(dose, unit, conc, dil, teor, fp, mult, equiv=1.0):
    """Replica o motor (importador PHP): UI/UFC/BLH via concentração; massa via mg. Sem ramo %."""
    unit=(unit or "").upper()
    teor = teor if teor>0 else 100.0
    if unit in ("UI","UFC","BLH"):
        if conc<=0: return None
        dose_ufc = dose*1e9 if unit=="BLH" else dose
        return (dose_ufc/conc)*equiv*dil/(teor/100.0)*fp*mult
    if unit in ("MG","MCG","G","",None) or unit=="ML":
        dose_mg = dose*1000 if unit=="G" else (dose/1000.0 if unit=="MCG" else dose)
        return (dose_mg*equiv*dil/(teor/100.0)*fp*mult)/1000.0
    return None

ap=argparse.ArgumentParser()
ap.add_argument("--db", required=True)
ap.add_argument("--commit", action="store_true")
args=ap.parse_args()
DRY = not args.commit

# ── Cadastro de ativos (mapas para re-match e recálculo) ─────────────────────
ativos = sb_all(f"/ativos?cliente_id=eq.{CID}&select=id,codigo_fc,nome,grupo,unidade_padrao,concentracao,teor,diluicao,fator_perda,densidade,preco_venda,custo_por_unidade,ativo")
by_id  = {a["id"]:a for a in ativos}
cod_by_id = {a["id"]:str(a.get("codigo_fc") or "").strip() for a in ativos}

def find_variante(nome, dilN):
    """Melhor ativo cujo nome casa (tokens, tolerante a espaço) E diluicao==dilN. Prefere unidade g e ativo=true."""
    tks=[t.upper() for t in toks(nome)]
    cand=[]
    for a in ativos:
        an=(a.get("nome") or "").upper()
        if abs(f(a.get("diluicao"),1.0)-dilN)>1e-6: continue
        if all(t in an for t in tks):
            cand.append(a)
    if not cand: return None
    cand.sort(key=lambda a:(0 if (a.get("unidade_padrao") or "").lower()=="g" else 1,
                            0 if a.get("ativo") else 1, str(a.get("nome"))))
    return cand[0]

# ── Aprovados com item diluído ("1:N") ───────────────────────────────────────
orcs = sb_all(f"/orcamentos?cliente_id=eq.{CID}&numero_orcamento=like.0001-*"
              "&select=id,numero_orcamento,status,forma_vol,qtde_potes,itens,valor_final_fc,total_orcamento")
orcs = [o for o in orcs if o.get("status") in APROVADOS]

# ── FCerta ───────────────────────────────────────────────────────────────────
con=fdb.connect(database=args.db, user="SYSDBA", password="masterkey", fb_library_name=FB_DLL, charset="NONE")

def fcerta_itens(nr, sr):
    c=con.cursor()
    c.execute("""SELECT CDPRO, DESCR, QUANT, QTREAL FROM FC12110
                 WHERE NRRQU=? AND SERIER=? AND TPCMP IN ('C','P')""",(nr,sr))
    out={}
    for cd,ds,qt,qr in c.fetchall():
        out[str(dec(cd)).strip()]={"descr":dec(ds),"qtreal":f(qr,None),"quant":f(qt,None)}
    return out   # sem close explícito (fdb 2.0.4 tropeça no reclose)
def fcerta_serier_raw(nr, sr_target):
    """SERIER é CHAR — retorna o char CRU do FC12100 cujo serier_int bate com a série do número."""
    c=con.cursor(); c.execute("SELECT SERIER FROM FC12100 WHERE NRRQU=?",(nr,))
    rows=[r[0] for r in c.fetchall()]
    for raw in rows:
        if serier_int(raw)==sr_target: return raw
    return rows[0] if rows else None

# ── Varredura ────────────────────────────────────────────────────────────────
linhas=[]   # planilha
plano=[]    # (orc, idx, novo_item)  p/ commit
om_plano=[] # (om_item_id, novo_qtd) p/ commit
n_orc_afetados=set()

for o in orcs:
    num=o["numero_orcamento"]; parts=num.split("-")
    if len(parts)!=3: continue
    try: nr=int(parts[1]); sr=serier_int(parts[2])
    except ValueError: continue
    itens=o.get("itens") or []
    if isinstance(itens,str): itens=json.loads(itens) or []
    mult=max(1.0,f(o.get("forma_vol"),1))*max(1,int(f(o.get("qtde_potes"),1)))
    fcit=None

    for idx,it in enumerate(itens):
        if it.get("tipo")!="mp" or it.get("is_qsp"): continue
        # "1:N" pode estar na prescrição OU no nome canônico (às vezes só um traz a diluição).
        nm_presc=it.get("nome_prescricao") or ""; nm_canon=it.get("nome") or ""
        m = DIL_RE.search(nm_presc) or DIL_RE.search(nm_canon)
        if not m: continue
        dilN=int(m.group(1))
        nome = nm_presc if DIL_RE.search(nm_presc) else nm_canon   # nome COM a diluição p/ re-match
        aid=it.get("ativo_id"); at_atual=by_id.get(aid)
        dil_atual=f((at_atual or {}).get("diluicao"), f(it.get("diluicao"),1.0))
        # re-match variante certa
        alvo=find_variante(nome, dilN)
        recasou = bool(alvo and at_atual and alvo["id"]!=at_atual["id"])
        base = alvo or at_atual
        if not base: continue
        conc=f(base.get("concentracao")); teor=f(base.get("teor"),100.0)
        dil =f(base.get("diluicao"),1.0); fp=f(base.get("fator_perda"),1.0) or 1.0
        preco=f(base.get("preco_venda"))
        upad=(base.get("unidade_padrao") or it.get("dose_unit") or "mg")
        dose=f(it.get("dose")); unit=it.get("dose_unit") or upad
        novo_g = recalc_qtd_g(dose, unit, conc, dil, teor, fp, mult)
        atual_g=f(it.get("qtd_total_g"))
        if novo_g is None:
            linhas.append([num,nome,str(cod_by_id.get(aid,"?")),str(base.get("codigo_fc")),
                           round(atual_g,4),"—","—","recálculo N/A","", "REVISAR"]);
            continue
        # FCerta ref pelo codigo do ativo CORRETO
        if fcit is None: fcit=fcerta_itens(nr, fcerta_serier_raw(nr, sr))
        ref=fcit.get(str(base.get("codigo_fc")).strip())
        qfc = (ref["qtreal"] if ref and ref["qtreal"] is not None else (ref["quant"] if ref else None)) if ref else None
        muda = abs(novo_g-atual_g) > max(1e-6, 1e-4*abs(atual_g))
        bate_fc = (qfc is not None) and abs(novo_g-qfc)<=max(0.001,0.01*abs(qfc))
        if not muda:
            continue  # já correto, ignora
        acao = "CORRIGIR" if bate_fc else ("REVISAR (≠FCerta)" if qfc is not None else "REVISAR (sem FCerta)")
        linhas.append([num,nome,
                       f"{cod_by_id.get(aid,'?')} (dil {dil_atual:g})",
                       f"{base.get('codigo_fc')} (dil {dil:g})"+(" ⟲recasa" if recasou else ""),
                       round(atual_g,4), round(novo_g,4),
                       (round(qfc,4) if qfc is not None else "—"),
                       ("sim" if bate_fc else "não"),
                       ("↑x%.3g"%(novo_g/atual_g) if atual_g else "—"), acao])
        if acao=="CORRIGIR":
            n_orc_afetados.add(num)
            novo_it=dict(it)
            novo_it["qtd_total_g"]=round(novo_g,4)
            if recasou:
                novo_it["ativo_id"]=base["id"]; novo_it["codigo_fc"]=str(base.get("codigo_fc") or "")
                novo_it["nome"]=(base.get("nome") or novo_it.get("nome"))
                novo_it["diluicao"]=dil; novo_it["teor"]=teor
                novo_it["densidade"]=f(base.get("densidade"),1.0) or 1.0
                novo_it["concentracao"]=conc; novo_it["preco_venda"]=preco
            # subtotal do item (custo) — proporcional à nova massa (não altera valor_final_fc travado)
            qtd_pad = round(novo_g*1000,4) if (upad or "mg").lower()!="g" else round(novo_g,4)
            novo_it["subtotal"]=round(qtd_pad*preco,4)
            plano.append((o, idx, novo_it, atual_g, novo_g, base))

try: con.close()
except Exception: pass

# ── OM: mapear qtd_pesar dos itens a corrigir (por orçamento) ────────────────
om_rows=[]
for o,idx,novo_it,atual_g,novo_g,base in plano:
    num=o["numero_orcamento"]
    oms=sbq(f"/lab_ordens?cliente_id=eq.{CID}&numero=eq.{num}&status=neq.cancelada&select=id&limit=1")
    if not oms:
        om_rows.append([num, base.get("nome"), "—", "sem OM"]); continue
    omid=oms[0]["id"]
    items=sbq(f"/lab_ordem_itens?ordem_id=eq.{omid}&select=id,ativo_id,descricao,qtd_pesar&order=ordem.asc")
    # casa SÓ por ativo_id exato (segurança: nunca adivinhar item de OM p/ não pesar o errado)
    alvo=None
    for li in items:
        if li.get("ativo_id") and li.get("ativo_id")==novo_it.get("ativo_id"): alvo=li; break
    if alvo:
        om_plano.append((alvo["id"], round(novo_g,4), num, base.get("nome")))
        om_rows.append([num, base.get("nome"), f'{f(alvo.get("qtd_pesar")):.4g} → {novo_g:.4g}', "ok"])
    else:
        om_rows.append([num, base.get("nome"), "—", "item da OM não localizado"])

# ── Planilha ─────────────────────────────────────────────────────────────────
HDR=Font(bold=True,color="FFFFFF"); FILLH=PatternFill("solid",fgColor="1F4E78"); CEN=Alignment(horizontal="center")
COR={"CORRIGIR":"C6EFCE","REVISAR (≠FCerta)":"FFC7CE","REVISAR (sem FCerta)":"FFEB9C","REVISAR":"FFEB9C"}
wb=openpyxl.Workbook()
ws=wb.active; ws.title="Correcoes"
ws.append(["Número","Item (prescrição)","Ativo ATUAL","Ativo CORRETO","qtd atual (g)","qtd recalc (g)","qtd FCerta (g)","bate FCerta?","fator","Ação"])
for r in linhas:
    ws.append(r)
    c=COR.get(r[9]);
    if c:
        for cell in ws[ws.max_row]: cell.fill=PatternFill("solid",fgColor=c)
for cell in ws[1]: cell.font=HDR; cell.fill=FILLH; cell.alignment=CEN
ws.freeze_panes="A2"
for col in ws.columns:
    w=max((len(str(c.value)) if c.value is not None else 0) for c in col)
    ws.column_dimensions[col[0].column_letter].width=min(max(w+2,10),46)

ws2=wb.create_sheet("OM_qtd_pesar")
ws2.append(["Número","Item","qtd_pesar (atual → novo)","Status"])
for r in om_rows: ws2.append(r)
for cell in ws2[1]: cell.font=HDR; cell.fill=FILLH; cell.alignment=CEN
ws2.freeze_panes="A2"
for col in ws2.columns:
    w=max((len(str(c.value)) if c.value is not None else 0) for c in col)
    ws2.column_dimensions[col[0].column_letter].width=min(max(w+2,10),44)

n_corrigir=sum(1 for r in linhas if r[9]=="CORRIGIR")
n_revisar =sum(1 for r in linhas if r[9].startswith("REVISAR"))
outx=r"C:\Users\carlo\Downloads\CORRECAO_DILUIDOS_2026-07-30.xlsx"
wb.save(outx)

print(f"Itens diluídos com divergência: {len(linhas)}  |  a CORRIGIR (bate FCerta): {n_corrigir}  |  REVISAR: {n_revisar}")
print(f"Orçamentos afetados: {len(n_orc_afetados)}  |  itens de OM a atualizar: {len(om_plano)}")
print("Planilha:", outx)

# ── GRAVAÇÃO (só --commit) ───────────────────────────────────────────────────
if DRY:
    print("\nDRY-RUN: nada gravado. Use --commit para aplicar as linhas 'CORRIGIR'.")
else:
    # agrupa por orçamento (um PATCH de itens por orçamento)
    por_orc={}
    for o,idx,novo_it,ag,ng,base in plano:
        por_orc.setdefault(o["id"], {"num":o["numero_orcamento"], "itens": (json.loads(o["itens"]) if isinstance(o["itens"],str) else list(o["itens"])) })
    for o,idx,novo_it,ag,ng,base in plano:
        por_orc[o["id"]]["itens"][idx]=novo_it
    err=0
    for oid,pk in por_orc.items():
        try: sb_patch(f"/orcamentos?id=eq.{oid}", {"itens": pk["itens"], "atualizado_em": datetime.datetime.utcnow().isoformat()+"Z"})
        except urllib.error.HTTPError as e: err+=1; print("  ERRO orc", pk["num"], e.read()[:120])
    for om_id, novo, num, nome in om_plano:
        try: sb_patch(f"/lab_ordem_itens?id=eq.{om_id}", {"qtd_pesar": novo})
        except urllib.error.HTTPError as e: err+=1; print("  ERRO OM", num, e.read()[:120])
    print(f"\nGRAVADO: {len(por_orc)} orçamentos + {len(om_plano)} itens de OM | erros {err}")
