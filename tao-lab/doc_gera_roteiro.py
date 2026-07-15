# -*- coding: utf-8 -*-
# Roteiro de Testes do TAO Lab em XLSX (fase, módulo, cenário, condição, passos, esperado, status).
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = r"C:\Users\carlo\Roteiro_Testes_TAO_Lab_v2.xlsx"

# (fase, modulo, cenario, precondicao, passos, esperado, prioridade)
T = [
 # FASE 0 — PREPARAÇÃO
 ("0. Preparação","Configurações","Preencher dados da farmácia","Usuário logado como administrador","Fórmulas→Configurações; preencher razão social, CNPJ, endereço, RT + CRF; Salvar","Dados salvos; passam a aparecer no rótulo","Alta"),
 ("0. Preparação","Configurações","Ligar Motor v2","Dados da farmácia salvos","Marcar 'Ativar no editor de orçamentos'; Salvar","Opção fica marcada após recarregar","Alta"),
 ("0. Preparação","Formas","Parametrizar validade por forma","Formas cadastradas","Cadastros→Formas; editar cada forma; informar 'Validade padrão (dias)'; Salvar","Prazo grava por forma e passa a valer na OM (não é fixo no código)","Alta"),
 ("0. Preparação","Formas","Parametrizar modo de preparo por forma","Formas cadastradas","Editar forma; preencher 'Modo de preparo / precauções'; Salvar","Texto grava; será herdado pela OM daquela forma","Média"),
 # FASE 1 — CADASTROS
 ("1. Cadastros","Ativos","Consultar ativo existente","Catálogo sincronizado do FCerta","Fórmulas→Ativos; buscar 'CAFEINA'; clicar no nome","Abre detalhe com preços, dados técnicos e sinônimos","Média"),
 ("1. Cadastros","Ativos","Criar ativo novo","—","+ Novo Ativo; preencher nome/grupo/unidades/preços; Salvar","Ativo aparece na lista e na busca","Alta"),
 ("1. Cadastros","Ativos","Editar dados técnicos","Ativo existente","Abrir ativo; Editar produto; alterar teor/diluição; Salvar","Alteração persiste ao reabrir","Média"),
 ("1. Cadastros","Prescritores","Cadastrar prescritor","—","Cadastros→Prescritores; + Novo; preencher CRM/UF/especialidade; Salvar","Prescritor listado e buscável","Média"),
 ("1. Cadastros","Prescritores","Endereço automático pelo CEP","Cadastro de prescritor aberto","No cadastro, digitar o CEP","Endereço/cidade/UF preenchidos sozinhos (ViaCEP)","Baixa"),
 ("1. Cadastros","Fornecedores","Cadastro único de fornecedor","—","Cadastros→Fornecedores (item único); + Novo; preencher CNPJ/IE/licenças; Salvar","Um só cadastro serve Fórmula e Cotações (mesma base)","Alta"),
 ("1. Cadastros","Cliente/Histórico","Editar dados do cliente","Cliente no histórico","Histórico; buscar cliente; Editar dados; marcar saúde; Salvar","Selo de saúde aparece; nome único no CRM","Média"),
 # FASE 2 — ORÇAMENTO
 ("2. Orçamento","Novo Orçamento","Montar fórmula simples","Motor v2 ligado","Novo Orçamento; paciente; forma Cápsula; add ativo + dose; QSP no excipiente; Salvar","Orçamento salvo; total calculado","Alta"),
 ("2. Orçamento","Novo Orçamento","Autocomplete de paciente","Base de contatos com clientes","Digitar o nome no campo Paciente; navegar ↑↓ e Enter","Busca na base do CRM; ao escolher preenche nome + WhatsApp","Média"),
 ("2. Orçamento","Novo Orçamento","Equivalência do sinônimo","Motor v2 ligado; sinônimo c/ fator","Buscar ativo por um sinônimo com equivalência (ex CISTEINA)","Quantidade ajustada pelo fator; dropdown mostra 'equiv'","Alta"),
 ("2. Orçamento","Novo Orçamento","Alerta de dose máxima","Ativo com dose_max","Digitar dose acima do máximo","Campo fica vermelho com tooltip de aviso (não bloqueia)","Alta"),
 ("2. Orçamento","Novo Orçamento","Trava de restrição","Ativo restricao=bloqueada ou SEMAGLUTIDA","Tentar adicionar o ativo bloqueado","Sistema recusa com alerta; item não entra","Alta"),
 ("2. Orçamento","Novo Orçamento","Fórmula padrão","Fórmulas padrão carregadas","Botão Fórmula padrão; buscar; aplicar","Linhas preenchidas com itens/doses da fórmula","Média"),
 ("2. Orçamento","Novo Orçamento","Análise de preços","Orçamento com itens","Abrir Análise de Preços; simular margem; aplicar","Valor final recalculado pela margem","Média"),
 # FASE 3 — HISTÓRICO
 ("3. Histórico","Histórico","Buscar e navegar por setas","Histórico carregado","Digitar nome; usar setas ↑↓ e Enter","Navega e abre o cliente selecionado","Baixa"),
 ("3. Histórico","Histórico","Ver resumo e componentes","Cliente com fórmulas","Expandir uma fórmula","Mostra resumo dos ativos + posologia","Média"),
 ("3. Histórico","Repetição","Repetir fórmula","Cliente com fórmula","Botão ↻ Repetir","Cai no editor com forma/cápsula/itens; obs com valor da época","Alta"),
 # FASE 4 — ESTOQUE ENTRADA NF
 ("4. Estoque","Entrada NF","Importar XML válido","Fornecedor com CNPJ cadastrado","Estoque→Entrada NF; Carregar XML da NF-e","Lê emitente, itens, lote/validade, duplicatas","Alta"),
 ("4. Estoque","Entrada NF","Associar item novo (de-para)","Item sem associação","Buscar e associar ao ativo","Associação memorizada p/ próxima NF do fornecedor","Alta"),
 ("4. Estoque","Entrada NF","Destino do valor","NF em conferência","Escolher Compra/Custo/Ambos num item; Efetivar","Preço do ativo atualiza conforme a escolha","Média"),
 ("4. Estoque","Entrada NF","Efetivar entrada","Todos itens associados","Clicar Efetivar","Cria lotes, movimenta estoque, gera contas a pagar","Alta"),
 ("4. Estoque","Entrada NF","NF duplicada","NF já efetivada","Reimportar a mesma NF","Sistema bloqueia (já importada)","Média"),
 # FASE 5 — LOTES
 ("5. Estoque","Lotes","Aprovar lote no CQ","Lote em quarentena (da NF)","Estoque→Lotes; ✔ Aprovar; informar laudo","Status muda p/ aprovado; registra quem/quando","Alta"),
 ("5. Estoque","Lotes","Reprovar lote","Lote em quarentena","✖ Reprovar; informar motivo","Status reprovado; não disponível p/ produção","Média"),
 ("5. Estoque","Lotes","Inventário/ajuste","Lote aprovado","Botão ⚖; informar qtd contada diferente","Gera movimento de ajuste; saldo corrigido","Média"),
 ("5. Estoque","Lotes","Kardex","Lote com movimentos","Botão ↔","Mostra extrato entradas/saídas","Baixa"),
 # FASE 6 — REPOSIÇÃO
 ("6. Estoque","Reposição","Definir mínimo","Ativo existente","Reposição; Definir mínimo; informar mín/máx/curva","Ativo passa a ser monitorado","Média"),
 ("6. Estoque","Reposição","Alerta abaixo do mínimo","Saldo < mínimo","Marcar 'só abaixo do mínimo'","Item destacado em vermelho com sugerido","Alta"),
 ("6. Estoque","Reposição","Gerar cotação","Itens abaixo do mínimo","Marcar itens; Gerar cotação","Cotação criada no módulo Cotações c/ Últ.Pago","Alta"),
 # FASE 7 — PRODUÇÃO
 ("7. Produção","Produção","Gerar OM do orçamento","Orçamento salvo","Produção; buscar orçamento; confirmar","OM criada na 1ª etapa; validade calculada","Alta"),
 ("7. Produção","Produção","Pesagem com lote (rastreabilidade)","OM aberta; lotes aprovados","Abrir OM; informar pesado + escolher lote FEFO","Só lotes aprovados; salva pesagem/lote","Alta"),
 ("7. Produção","Produção","Mover no kanban","OM aberta","Mover para próxima etapa","OM muda de coluna; registra auditoria","Média"),
 ("7. Produção","Produção","Concluir + baixar estoque","OM pesada; etapa final","Mover p/ etapa final; confirmar","OM concluída; estoque dos lotes baixado (kardex)","Alta"),
 ("7. Produção","Produção","Baixa idempotente","OM já concluída","Tentar concluir de novo","Não baixa estoque em dobro","Média"),
 ("7. Produção","Validade (RDC 67)","Validade reduzida pelo lote","OM pesada; lote vence antes do prazo da forma","Concluir a OM","Validade final = MENOR entre o prazo da forma e a validade do lote; alerta exibido; rótulo com a data correta","Alta"),
 ("7. Produção","Controlados (344)","Trava de controlado sem receita","OM com componente controlado, sem dados da receita","Tentar concluir a OM","Bloqueia; exige tipo de receita, nº notificação, comprador e prescritor","Alta"),
 ("7. Produção","Controlados (344)","Concluir controlado com receita","OM controlada","Preencher o bloco '🔒 Receita controlada'; concluir","Libera; gera a saída no SNGPC","Alta"),
 ("7. Produção","Modo de preparo","Modo de preparo na OM","Forma com modo de preparo cadastrado","Abrir OM; conferir Modo de preparo (herdado da forma); ajustar; Salvar","Texto aparece herdado; ajuste específico persiste na OM","Média"),
 # FASE 8 — RÓTULO + LIVRO
 ("8. Rótulo/Livro","Rótulo","Emitir rótulo RDC 67","OM criada; dados da farmácia preenchidos","Botão 🏷 Rótulo na OM","Janela imprimível com todos os dizeres RDC 67","Alta"),
 ("8. Rótulo/Livro","Rótulo","Aviso dados faltando","empresa_config vazio","Emitir rótulo sem dados da farmácia","Mostra aviso p/ preencher Configurações","Média"),
 ("8. Rótulo/Livro","Livro Receituário","Gerar livro por período","OMs no período","Livro de Receituário; escolher período; Gerar","Lista sequencial de OMs; imprimível","Alta"),
 # FASE 9 — CONTAS A PAGAR
 ("9. Financeiro","Contas a Pagar","Ver duplicatas da NF","NF efetivada com duplicatas","Fórmulas→Contas a Pagar","Duplicatas listadas em aberto","Alta"),
 ("9. Financeiro","Contas a Pagar","Marcar como paga","Conta em aberto","Botão ✔ pagar","Status pago com data; total atualiza","Média"),
 ("9. Financeiro","Contas a Pagar","Relatório ao contador","Contas cadastradas","Botão 🖨 Relatório","Versão imprimível sem botões","Baixa"),
 # FASE 11 — FORNECEDORES (detalhado / RDC 67)
 ("11. Fornecedores","Fornecedor","Cadastro fiscal completo","—","Fórmulas→Fornecedores; + Novo; preencher CNPJ, IE, endereço, tipo, regime (CRT)","Fornecedor salvo com dados fiscais p/ casar NF","Alta"),
 ("11. Fornecedores","Fornecedor","Licenças + qualificação (RDC 67)","Fornecedor existente","Editar; preencher AFE/AE/Licença VISA c/ validade; marcar Qualificado","Badges de licença aparecem na lista; vencidas ficam em vermelho","Média"),
 ("11. Fornecedores","Integração NF","Cadastro rápido pelo XML","NF de fornecedor não cadastrado","Estoque→Entrada NF; carregar XML; botão 'Cadastrar fornecedor com dados da NF'","Fornecedor criado já preenchido (razão, IE, endereço); Efetivar libera","Alta"),
 ("11. Fornecedores","Integração NF","Alerta de licença vencida","Fornecedor c/ AFE/licença vencida","Carregar XML desse fornecedor","Banner de conformidade RDC 67 (⛔ vencida / ⚠ não qualificado); não bloqueia","Média"),
 # FASE 12 — FICHA TÉCNICA + LAUDO (RDC 67)
 ("12. RDC 67","Ficha técnica","Preencher ficha da MP","Ativo (matéria-prima)","Ativos; abrir ativo; Editar; seção Ficha técnica; preencher fórmula/PM/pF/pH/solubilidade","Ficha salva; reabre preenchida","Média"),
 ("12. RDC 67","Laudo por lote","Anexar laudo de análise","Lote existente","Estoque→Lotes; botão 📎 Laudo; nº do certificado + PDF; Salvar","Ícone vira 📄; abre o laudo salvo (uploads/laudos-mp)","Alta"),
 # FASE 13 — INVENTÁRIO EM MASSA
 ("13. Inventário","Inventário","Abrir sessão (congela saldo)","Lotes com saldo","Estoque→Inventário; + Novo; escopo aprovados; Abrir","Planilha de contagem com o saldo congelado","Alta"),
 ("13. Inventário","Inventário","Contagem + divergência","Sessão aberta","Digitar qtd contada diferente do sistema","Diferença aparece (verde/vermelho); salva sozinha","Média"),
 ("13. Inventário","Inventário","Fechar e aplicar ajustes","Sessão com divergências","Fechar inventário","Ajustes aplicados em lote no kardex; não contados intactos","Alta"),
 # FASE 14 — HISTÓRICO DE PREÇOS
 ("14. Preços","Histórico","Ver linha do tempo","Ativo com NF/edição de preço","Ativos; abrir ativo; 📈 Histórico de preços","Lista datada de compra/custo/venda por origem","Baixa"),
 # FASE 15 — PRODUÇÃO INTERNA (diluições/bases)
 ("15. Prod. Interna","Produção Interna","Nova produção de diluição","Ativo diluído c/ receita (fórmula padrão)","Produção Interna; + Nova; escolher diluído + quantidade","Receita escalada; insumos calculados (puro + veículo)","Alta"),
 ("15. Prod. Interna","Produção Interna","Pesagem FEFO + gerar lote","Produção aberta; insumos c/ lote","Pesar cada insumo (lote FEFO); Concluir","Baixa insumos; gera lote PI-AAAAMM-NNN do diluído (teor/fator/lote_puro)","Alta"),
 # FASE 16 — CONTROLADOS / SNGPC (Portaria 344/98)
 ("16. Controlados","Ativos","Marcar ativo controlado","Ativo de substância controlada","Editar ativo; marcar Controlado; classe SNGPC (A1/B1/C1...)","Ativo passa a escriturar no SNGPC","Alta"),
 ("16. Controlados","SNGPC","Lançamento manual","Ativo controlado","Controlados/SNGPC; lançar entrada/saída/perda","Movimento registrado no livro","Alta"),
 ("16. Controlados","SNGPC","Escrituração automática da OM","OM controlada concluída","Concluir OM com componente controlado","Gera saída SNGPC automática (idempotente)","Alta"),
 ("16. Controlados","SNGPC","Balanço + XML","Movimentos no período","Balanço BMPO; Gerar XML","Saldo por substância; XML p/ envio à ANVISA","Média"),
 # FASE 17 — LGPD + CID
 ("17. LGPD/CID","LGPD","Consentimento do paciente","Cliente no cadastro","Histórico; editar cliente; marcar Consentimento + data + canal; Salvar","Consentimento registrado; acesso fica na trilha","Média"),
 ("17. LGPD/CID","LGPD","Anonimizar paciente","Cliente existente (perfil gestor)","Editar cliente; 🔒 Anonimizar; confirmar","Nome/dados apagados; OMs preservadas por rastreabilidade","Baixa"),
 ("17. LGPD/CID","CID-10","Diagnóstico opcional no orçamento","Catálogo CID carregado","Novo Orçamento; campo CID; buscar 'diabetes'","Autocomplete E10.x; grava no orçamento; NUNCA obrigatório","Baixa"),
 # FASE 18 — ENTREGAS
 ("18. Entregas","Entrega","Gatilho na aprovação","Card com pedido a entregar","Fechar o card como ganho (vai p/ Pós-vendas)","Entrega criada automaticamente (status pendente), vinculada ao cliente","Alta"),
 ("18. Entregas","Entrega","Endereço pelo CEP","Card em Pós-vendas; aba Entrega","Digitar o CEP ou escolher endereço salvo do cliente","Endereço preenchido (ViaCEP); pode salvar novo endereço p/ o cliente","Média"),
 ("18. Entregas","Kanban","Trava de endereço","Entrega motoboy/correio sem endereço","Arrastar o card p/ 'Pronto para Entrega'","Bloqueia e pede o endereço; tipo Balcão passa direto","Alta"),
 ("18. Entregas","Kanban","Trava de pagamento","Entrega ainda não paga","Arrastar o card p/ 'NPS'","Bloqueia até o pagamento estar registrado","Alta"),
 # FASE 19 — CAIXA
 ("19. Caixa","Vendas","Baixa manual do pagamento","Card ganho (venda 'a receber')","Financeiro→Caixa→Vendas; Receber; forma + parcelas/split; confirmar","Venda quitada; recibo e pagamento gerados (com taxa)","Alta"),
 ("19. Caixa","Entrega↔Caixa","Pagar na entrega baixa no Caixa","Card com entrega e venda aberta","Aba Entrega; marcar pago + forma; Salvar","Venda do card quitada no Caixa automaticamente","Alta"),
 ("19. Caixa","Entrega↔Caixa","Baixa no Caixa libera NPS","Venda do card quitada na tela de Vendas","Dar baixa no Caixa","Entrega marca paga; card pode ir p/ NPS","Média"),
 ("19. Caixa","Estorno","Estorno auditado","Venda quitada","Botão Estornar; informar o motivo","Recibo estornado (motivo/quem/quando); venda reaberta; pagamento estornado","Média"),
 # FASE 20 — COTAÇÕES
 ("20. Cotações","Nova Cotação","Fornecedores em chips + revisão","Itens e fornecedores (chips) selecionados","'Revisar mensagem e enviar'; editar o texto ({fornecedor}); confirmar","Mensagem revisável pelo farmacêutico; só envia após confirmar","Alta"),
 ("20. Cotações","Proposta","Upload de proposta por IA","Cotação criada, com fornecedores","Botão 📎 Proposta no fornecedor; subir PDF/foto","IA extrai os preços; comparativo atualizado","Média"),
 # FASE 10 — FLUXO E2E
 ("10. Fluxo E2E","Ponta a ponta","Ciclo completo da fórmula","Todos os cadastros e config prontos","Orçamento → (Repetir opcional) → gerar OM → pesar c/ lote → concluir (baixa estoque) → imprimir rótulo → conferir no Livro","Fórmula percorre todo o fluxo sem erro; rastreabilidade lote→OM→paciente íntegra","Alta"),
 ("10. Fluxo E2E","Ponta a ponta","Ciclo de compra","Fornecedor c/ CNPJ; ativo c/ mínimo","Reposição gera cotação → (compra) → Entrada NF → lote → CQ aprova → saldo reposto","Estoque reposto e rastreável; contas a pagar geradas","Alta"),
]

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Roteiro de Testes"
COLS = ["ID","Fase","Módulo","Cenário","Pré-condição","Passos","Resultado Esperado","Prioridade","Status","Data teste","Observações"]
azul = PatternFill("solid", fgColor="1E40AF"); branco = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="D0D7DE"); bord = Border(left=thin,right=thin,top=thin,bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
fases_fill = {}
palette = ["EFF6FF","F0FDF4","FEF3C7","FCE7F3","F3E8FF","ECFEFF","FEF2F2","F1F5F9","FFF7ED","E0F2FE","DCFCE7"]

# cabeçalho
ws.append(COLS)
for c in ws[1]:
    c.fill = azul; c.font = branco; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = bord
ws.row_dimensions[1].height = 24

fi = 0; last_fase = None
for i,(fase,mod,cen,pre,pas,esp,prio) in enumerate(T, start=1):
    if fase != last_fase: fi += 1; last_fase = fase
    fill = PatternFill("solid", fgColor=palette[(fi-1) % len(palette)])
    row = [f"T{i:02d}", fase, mod, cen, pre, pas, esp, prio, "", "", ""]
    ws.append(row)
    for c in ws[i+1]:
        c.alignment = wrap; c.border = bord; c.fill = fill

# larguras
larg = {"A":6,"B":14,"C":16,"D":30,"E":26,"F":42,"G":40,"H":10,"I":12,"J":12,"K":26}
for col,w in larg.items(): ws.column_dimensions[col].width = w

# validação de dados na coluna Status
dv = DataValidation(type="list", formula1='"OK,NOK,Bloqueado,N/A"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"I2:I{len(T)+1}")
dvp = DataValidation(type="list", formula1='"Alta,Média,Baixa"', allow_blank=True)
ws.add_data_validation(dvp); dvp.add(f"H2:H{len(T)+1}")

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:K{len(T)+1}"

# aba de resumo
ws2 = wb.create_sheet("Resumo")
ws2["A1"] = "Resumo por fase"; ws2["A1"].font = Font(bold=True, size=13, color="1E40AF")
ws2.append([]); ws2.append(["Fase","Total de casos","OK","NOK","Pendentes"])
for c in ws2[3]: c.font = Font(bold=True)
fases = []
for fase,*_ in T:
    if fase not in fases: fases.append(fase)
r0 = 4
for k,fase in enumerate(fases):
    n = sum(1 for f,*_ in T if f == fase)
    rr = r0 + k
    ws2.append([fase, n,
        f'=COUNTIFS(\'Roteiro de Testes\'!B:B,A{rr},\'Roteiro de Testes\'!I:I,"OK")',
        f'=COUNTIFS(\'Roteiro de Testes\'!B:B,A{rr},\'Roteiro de Testes\'!I:I,"NOK")',
        f'=B{rr}-C{rr}-D{rr}'])
ws2.append(["TOTAL", len(T),
    '=COUNTIF(\'Roteiro de Testes\'!I:I,"OK")',
    '=COUNTIF(\'Roteiro de Testes\'!I:I,"NOK")',
    f'=B{r0+len(fases)}-C{r0+len(fases)}-D{r0+len(fases)}'])
for col,w in {"A":16,"B":14,"C":8,"D":8,"E":12}.items(): ws2.column_dimensions[col].width = w

wb.save(OUT)
print("XLSX gerado:", OUT, "-", len(T), "casos de teste")
