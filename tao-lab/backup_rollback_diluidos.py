#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
"""
BACKUP + ROLLBACK das correções de pesagem de diluídos (30/07).
Lê a planilha CORRECAO_DILUIDOS_2026-07-30.xlsx (que tem os valores ANTIGOS na coluna
"qtd atual" e "OM atual→novo"), re-localiza os IDs no Supabase e gera:
  • backup_diluidos_2026-07-30.json  — snapshot ATUAL (itens completos dos 10 orçamentos + 8 OM itens)
  • rollback_diluidos_2026-07-30.json — mapa p/ REVERTER (qtd_total_g e qtd_pesar antigos)
--apply  reverte (aplica os valores antigos do rollback já gravado).
"""
import json, urllib.request, urllib.error, argparse, re, os
import openpyxl

SB="https://gclayesytzzpzkjvgede.supabase.co/rest/v1"; KEY=os.environ.get("SUPABASE_KEY", "sb_secret_HpoqM6ujk2yD6la7KM3cuQ_pdWBK8jo")
CID="62f98634-77ff-42f4-acaf-8561d56583da"
XLSX=r"C:\Users\carlo\Downloads\CORRECAO_DILUIDOS_2026-07-30.xlsx"
BK=r"C:\Users\carlo\Downloads\backup_diluidos_2026-07-30.json"
RB=r"C:\Users\carlo\Downloads\rollback_diluidos_2026-07-30.json"

def sbq(p):
    h={"apikey":KEY,"Authorization":"Bearer "+KEY}
    return json.loads(urllib.request.urlopen(urllib.request.Request(SB+p,headers=h)).read().decode())
def sb_patch(path, body):
    r=urllib.request.Request(SB+path, data=json.dumps(body).encode(), method="PATCH",
        headers={"apikey":KEY,"Authorization":"Bearer "+KEY,"Content-Type":"application/json","Prefer":"return=minimal"})
    return urllib.request.urlopen(r).status
def num_of(v):
    try: return float(str(v).replace(",","."))
    except (TypeError,ValueError): return None

ap=argparse.ArgumentParser(); ap.add_argument("--apply",action="store_true"); args=ap.parse_args()

if args.apply:
    rb=json.load(open(RB,encoding="utf-8"))
    err=0
    for o in rb["orcamentos"]:
        cur=sbq(f"/orcamentos?id=eq.{o['orc_id']}&select=itens")[0]["itens"]
        cur[o["idx"]]["qtd_total_g"]=o["qtd_total_g_old"]
        if o.get("subtotal_old") is not None: cur[o["idx"]]["subtotal"]=o["subtotal_old"]
        try: sb_patch(f"/orcamentos?id=eq.{o['orc_id']}", {"itens":cur})
        except urllib.error.HTTPError as e: err+=1; print("ERRO orc",o["num"],e.read()[:100])
    for m in rb["om_itens"]:
        try: sb_patch(f"/lab_ordem_itens?id=eq.{m['om_item_id']}", {"qtd_pesar":m["qtd_pesar_old"]})
        except urllib.error.HTTPError as e: err+=1; print("ERRO OM",m["num"],e.read()[:100])
    print(f"ROLLBACK aplicado: {len(rb['orcamentos'])} orçamentos + {len(rb['om_itens'])} OM itens | erros {err}")
    raise SystemExit

wb=openpyxl.load_workbook(XLSX)
# CORRIGIR: num -> qtd_antigo
corr={}
for r in wb["Correcoes"].iter_rows(min_row=2, values_only=True):
    num,item,ata,atc,qa,qn,qfc,bate,fator,acao=r
    if acao=="CORRIGIR": corr[num]={"qtd_old":num_of(qa),"qtd_new":num_of(qn)}
# OM: num -> qtd_pesar antigo (parse "X -> Y" / "X → Y")
om_old={}
for r in wb["OM_qtd_pesar"].iter_rows(min_row=2, values_only=True):
    num,item,trans,status=r
    m=re.match(r"\s*([\d.,]+)\s*[-→>]+\s*([\d.,]+)", str(trans))
    if m and status=="ok": om_old[num]=num_of(m.group(1))

backup={"orcamentos":[], "om_itens":[]}; rollback={"orcamentos":[], "om_itens":[]}
for num,c in corr.items():
    o=sbq(f"/orcamentos?cliente_id=eq.{CID}&numero_orcamento=eq.{num}&select=id,itens")
    if not o: print("!! orçamento não achado", num); continue
    oid=o[0]["id"]; itens=o[0]["itens"]
    idx=None
    for i,it in enumerate(itens):
        if it.get("tipo")=="mp" and "VIT D3" in (it.get("nome") or "").upper(): idx=i; break
    if idx is None: print("!! item VIT D3 não achado", num); continue
    it=itens[idx]
    backup["orcamentos"].append({"num":num,"orc_id":oid,"idx":idx,"item_atual":it})
    rollback["orcamentos"].append({"num":num,"orc_id":oid,"idx":idx,"ativo_id":it.get("ativo_id"),
        "qtd_total_g_old":c["qtd_old"], "qtd_total_g_new_confirmado":it.get("qtd_total_g")})
    # OM
    if num in om_old:
        oms=sbq(f"/lab_ordens?cliente_id=eq.{CID}&numero=eq.{num}&status=neq.cancelada&select=id&limit=1")
        if oms:
            for li in sbq(f"/lab_ordem_itens?ordem_id=eq.{oms[0]['id']}&select=id,ativo_id,descricao,qtd_pesar"):
                if li.get("ativo_id")==it.get("ativo_id"):
                    backup["om_itens"].append({"num":num,"om_item_id":li["id"],"qtd_pesar_atual":li.get("qtd_pesar")})
                    rollback["om_itens"].append({"num":num,"om_item_id":li["id"],"qtd_pesar_old":om_old[num],
                        "qtd_pesar_new_confirmado":li.get("qtd_pesar")}); break

json.dump(backup, open(BK,"w",encoding="utf-8"), ensure_ascii=False, indent=2)
json.dump(rollback, open(RB,"w",encoding="utf-8"), ensure_ascii=False, indent=2)
print(f"Snapshot atual : {len(backup['orcamentos'])} orçamentos + {len(backup['om_itens'])} OM itens -> {BK}")
print(f"Mapa rollback  : {len(rollback['orcamentos'])} orçamentos + {len(rollback['om_itens'])} OM itens -> {RB}")
print("Reverter tudo:  python backup_rollback_diluidos.py --apply")
