# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REAL = r"C:\Users\carlo\tao-crm\_manuais\cot"
base = r"C:\Users\carlo\OneDrive\Área de Trabalho\Soluções & TAO\Manuais"
if not os.path.isdir(base): base = REAL
OUT  = os.environ.get("COTMAN_OUT", os.path.join(base, "Manual_TAO_Cotacoes.docx"))

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.9); s.right_margin = Inches(0.9)
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)

_ch=[0]; _sec=[0]
def H1(t):
    _ch[0]+=1; _sec[0]=0; doc.add_heading(f"{_ch[0]}. {t}", level=1)
def H2(t):
    _sec[0]+=1; doc.add_heading(f"{_ch[0]}.{_sec[0]} {t}", level=2)
def H3(t): doc.add_heading(t, level=3)
def P(t=""): doc.add_paragraph(t)
def B(t): doc.add_paragraph(t, style="List Bullet")
def N(t): doc.add_paragraph(t, style="List Number")
def PB(): doc.add_page_break()
def CAP(t):
    cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=cp.add_run(t); r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0x66,0x66,0x66)
def IMG(fname, caption=None, w=6.3):
    path=os.path.join(REAL, fname)
    if os.path.exists(path):
        pp=doc.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        try: pp.add_run().add_picture(path, width=Inches(w))
        except Exception as e: P("[img falhou: %s]"%fname)
        if caption: CAP(caption)
    else:
        P("[imagem pendente: "+fname+"]")
def TOC():
    p=doc.add_paragraph(); r=p.add_run()._r
    b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin'); r.append(b)
    i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text='TOC \\o "1-2" \\h \\z \\u'; r.append(i)
    s=OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'),'separate'); r.append(s)
    t=OxmlElement('w:t'); t.text="Abra no Word e clique com o botão direito > Atualizar campo."; r.append(t)
    e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end'); r.append(e)

# ── CAPA ──
tp=doc.add_paragraph(); tp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=tp.add_run("TAO Neo"); r.bold=True; r.font.size=Pt(40); r.font.color.rgb=RGBColor(0x15,0x2C,0x42)
st=doc.add_paragraph(); st.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=st.add_run("Manual de Operação — Módulo de Cotações"); r.font.size=Pt(20); r.font.color.rgb=RGBColor(0x55,0x60,0x70)
sd=doc.add_paragraph(); sd.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=sd.add_run("Do cadastro de fornecedores à sugestão de pedido"); r.italic=True; r.font.size=Pt(12)
P(); P()
P("Este manual mostra como operar o módulo de Cotações de compra de insumos: cadastrar fornecedores, montar uma cotação e priorizar itens, enviar a solicitação, registrar as propostas recebidas, conferir a associação dos itens, comparar preços (com frete) e gerar a sugestão de pedido por fornecedor. As telas são reais do sistema.")
PB()
doc.add_heading("Sumário", level=1); TOC()

# 1 Visão geral
PB(); H1("Visão geral e fluxo")
P("O módulo de Cotações organiza a compra de insumos de ponta a ponta, sempre dentro do TAO Neo. O fluxo é:")
N("Cadastrar os fornecedores (uma vez; reaproveitados em todas as cotações).")
N("Criar uma cotação com os itens a comprar, definindo a prioridade de cada um e os fornecedores participantes.")
N("Enviar a solicitação aos fornecedores pelo WhatsApp (com revisão do texto antes do envio).")
N("Registrar o retorno de cada fornecedor (PDF/foto lido por IA ou por um modelo aprendido, ou digitação manual).")
N("Conferir a associação de cada item retornado ao ativo do TAO Neo (farmacêutico).")
N("Analisar o comparativo de preços (normalizado, com frete) e escolher o fornecedor.")
N("Gerar a sugestão de pedido por fornecedor (por preço ou consolidado) e ajustar conforme a necessidade.")
P("Onde encontrar: menu Cotações (com os submenus Cotações, Nova Cotação, Fornecedores e Modelos de Proposta).")
IMG("real_cot_lista.png", "Cotações — lista das cotações.")

# 2 Fornecedores
PB(); H1("Cadastro de Fornecedores")
P("O fornecedor é um cadastro único, reaproveitado em todas as cotações. Vá em Cotações > Fornecedores.")
IMG("real_cot_forn_lista.png", "Cotações > Fornecedores — lista.")
H2("Cadastrar / editar um fornecedor")
P("Clique em “+ Novo Fornecedor” (ou no lápis para editar). Preencha os dados:")
B("Nome e WhatsApp (com DDD, só números) — o WhatsApp é o que permite enviar a cotação e receber a proposta pelo próprio módulo.")
B("Pessoa de contato, CNPJ (casa com o XML da NF), razão social, endereço, e-mail e telefone.")
B("Prazo de pagamento.")
B("Pedido mínimo (R$) — valor mínimo de pedido/faturamento do fornecedor. É usado na sugestão de pedido: quando o pedido sugerido daquele fornecedor fica abaixo desse valor, o sistema exibe um alerta (não impede a compra).")
IMG("real_cot_forn_form.png", "Cadastro de fornecedor — com o campo Pedido mínimo.")
P("Observação: fornecedor já usado em alguma cotação é desativado em vez de excluído, para preservar o histórico.")

# 3 Nova cotação
PB(); H1("Criar uma Cotação")
P("Em Cotações > Nova Cotação você monta a cotação em três blocos: itens, fornecedores e envio.")
IMG("real_cot_nova.png", "Nova Cotação — itens, fornecedores e identificação.")
H2("1. Itens para cotação")
P("Adicione os ativos que quer cotar. Você pode importar de uma planilha (com código, descrição, unidade e quantidade sugerida) ou incluir item a item buscando o ativo do cadastro. Marque pelo menos um item como urgente (⭐) — os urgentes são os mandatórios da compra.")
H2("2. Fornecedores participantes")
P("Selecione os fornecedores que vão receber a cotação. O sistema sugere os mais frequentes. É possível navegar e escolher pelo teclado (setas ↑/↓ e Enter).")
H2("3. Identificação e envio")
P("Dê um título à cotação e escolha a instância de WhatsApp de envio. Ao salvar, a cotação é criada e você é levado à tela de detalhe, de onde envia e registra os retornos.")

# 4 Prioridades
PB(); H1("Prioridade dos itens")
P("Na tela de detalhe da cotação, a seção Itens lista tudo que será cotado. Cada item tem um seletor de PRIORIDADE, que organiza a distribuição na sugestão de pedido:")
B("0 – Urgente (⭐): item mandatório; equivale ao “favorito”. Determina os fornecedores principais na visão Consolidado.")
B("1 – 15 dias / 2 – 30 dias / 3 – Acima de 30 dias: urgência decrescente.")
IMG("real_cot_itens.png", "Seção Itens — prioridade, edição e seleção múltipla.")
H2("Editar, incluir e excluir itens")
P("Na própria lista você pode:")
B("Alterar qualquer item: descrição, código, quantidade, unidade e prioridade (salva automaticamente).")
B("Incluir um novo item pela linha do rodapé (com busca de ativo pelo teclado).")
B("Excluir um item (🗑) ou vários de uma vez: marque as caixas de seleção e use “Excluir selecionados”.")
H2("Conferência do farmacêutico")
P("O botão “🔍 Conferência do farmacêutico” (no topo da seção Itens) abre a revisão de todos os itens retornados e sua associação ao ativo do TAO Neo — ver o capítulo de Conferência.")

# 5 Enviar
PB(); H1("Enviar aos Fornecedores")
P("Na tela de detalhe, clique em “📤 Enviar aos fornecedores”. Abre a tela de revisão do texto:")
B("O texto que será enviado aparece editável. O marcador {fornecedor} é trocado pelo nome de cada fornecedor no envio.")
B("A lista de destinatários mostra quem tem WhatsApp e o status (já enviado / respondeu).")
P("O envio só acontece quando você clica em “✅ Confirmar e enviar”. Alternativamente:")
B("“💾 Salvar (não enviar)”: guarda o texto revisado sem disparar nada.")
B("“📋 Copiar texto”: copia para envio manual (útil para fornecedor sem WhatsApp).")
P("Observação: o texto lista apenas os itens (sem quantidades), com ⭐ nos prioritários.")
IMG("real_cot_envio.png", "Enviar aos fornecedores — revisão do texto antes do envio.")

# 6 Registrar retorno
PB(); H1("Registrar o Retorno do Fornecedor")
P("Quando o fornecedor responde, clique em “📥 Registrar retorno de fornecedor”, escolha o fornecedor e informe a proposta. Não é preciso ter enviado pelo módulo — vale para cotação feita por fora.")
IMG("real_cot_registrar_retorno.png", "Registrar retorno — escolha do fornecedor e origem.")
H2("Como o sistema lê a proposta")
P("Há dois caminhos automáticos e um manual:")
B("Modelo aprendido (sem IA): se o fornecedor já tem um layout aprendido (ver Modelos de Proposta), o PDF é lido de forma determinística — rápido e sem custo.")
B("IA (Gemini): sem modelo, a IA lê o PDF/foto e extrai os itens, preços, fracionamento e validade.")
B("Manual: digitar a proposta item a item.")
P("O sistema normaliza os preços para uma base comparável (R$/g, R$/ml ou R$/milheiro) e converte o fracionamento mínimo informado pelo fornecedor.")
H2("Validação antes de gravar")
P("Antes de efetivar, o farmacêutico vê a tela de revisão dos itens lidos e pode ajustar preço, unidade, fracionamento, validade e a associação ao ativo. Só ao confirmar os dados vão para o comparativo. Ao registrar o retorno, a cotação passa para “recebendo”.")

# 7 Modelos de proposta
PB(); H1("Modelos de Proposta (aprendizado de layout)")
P("Para não usar IA a cada importação, o módulo aprende o layout de cada fornecedor uma vez e reaproveita nas próximas. Em Cotações > Modelos de Proposta você vê os modelos existentes e pode criar um novo:")
N("Suba um PDF de exemplo do fornecedor.")
N("Clique em “🤖 Analisar (IA propõe o layout)” — a IA identifica as colunas (item, preço, unidade, fracionamento, validade).")
N("Revise e salve. A partir daí, as próximas propostas daquele fornecedor são lidas sem IA.")
IMG("real_cot_modelos.png", "Modelos de Proposta — layout aprendido por fornecedor.")

# 8 Conferência
PB(); H1("Conferência do Farmacêutico")
P("A conferência é a revisão da associação: garantir que cada item retornado aponta para o ativo correto do TAO Neo. Abra pelo botão “🔍 Conferência do farmacêutico” na seção Itens. É opcional (não trava o processo) e cada associação já é salva na hora.")
IMG("real_cot_conferencia.png", "Conferência do farmacêutico — status e associação por item.")
P("Cada linha mostra o status:")
B("🟢 casado: associado a um item da cotação.")
B("🟡 fora da lista: associado a um ativo que não está entre os itens desta cotação (confira se está certo).")
B("🔴 não associado: precisa associar.")
P("Para corrigir, escolha o ativo no campo de busca (navega pelo teclado). A associação vira sinônimo — as próximas cotações casam sozinhas. Também dá para excluir um preço e incluir uma nova linha. A tela rola verticalmente; feche por “Salvar e fechar” ou pelo ✕ (clicar fora não fecha).")

# 9 Comparativo
PB(); H1("Comparativo de Preços")
P("O comparativo mostra, por item, o preço de cada fornecedor já normalizado, destacando o melhor. É onde se decide de quem comprar.")
IMG("real_cot_comparativo.png", "Comparativo — melhor fornecedor, frete e seleção.")
H2("Leitura das cores e do melhor")
B("Verde: melhor preço do item.")
B("Vermelho: melhor preço acima do último valor pago.")
H2("Frete (rateio por valor)")
P("No topo de cada fornecedor há “✎ retorno”, que abre a edição daquela proposta (preço, fracionamento, validade) e o campo de FRETE. O frete é rateado proporcionalmente ao valor entre os itens do fornecedor; o comparativo passa a escolher o melhor pelo preço COM frete e mostra também o valor SEM frete ao lado.")
IMG("real_cot_editret.png", "Editar retorno do fornecedor — proposta e campo de Frete.")
H2("Trocar o fornecedor escolhido")
P("Na coluna MELHOR FORN. há um seletor por item: fica em “Auto” (o menor preço sugerido pelo sistema) e você pode fixar outro fornecedor que cotou aquele item. A escolha é salva e reflete na sugestão de pedido.")
B("“↺ Restaurar sugestões”: volta todos os itens ao automático.")
B("Para reverter só um item: escolha a opção “Auto” no seletor dele.")
P("Use “⬇️ Exportar XLSX” para baixar o comparativo completo.")

# 10 Sugestão de pedido
PB(); H1("Sugestão de Pedido")
P("Abaixo do comparativo, a seção “🧾 Sugestão de pedido” monta o pedido por fornecedor. O melhor preço é o critério; a prioridade organiza a distribuição.")
IMG("real_cot_sugestao.png", "Sugestão de pedido — por fornecedor, com frete e alertas.")
H2("As duas visões")
B("Por preço: cada item vai para o fornecedor mais barato.")
B("Consolidado: os favoritos (⭐) definem os fornecedores principais e puxam os demais itens para eles (o mais barato entre os principais; se nenhum principal cotou, o mais barato geral).")
P("Alterne entre as visões nos botões. A escolha manual feita no comparativo é sempre respeitada nas duas.")
H2("Colunas e ajustes")
P("Por fornecedor, a tabela traz: produto (com a prioridade), quantidade necessária, quantidade de compra (arredondada ao fracionamento mínimo e editável), valor unitário sem frete, frete do item (rateado por valor), valor unitário com frete e valor total. Há o subtotal do fornecedor e o total geral do pedido.")
B("Mover um item de fornecedor: use o seletor na linha (só aparecem os fornecedores que cotaram) — os valores recalculam ao vivo.")
B("Ajustar a quantidade de compra: edite direto na célula.")
B("Alerta de pedido mínimo: quando o valor de produtos de um fornecedor fica abaixo do mínimo dele, aparece um aviso com quanto falta (não impede o pedido).")

# 11 Dicas
PB(); H1("Dicas e observações")
B("Combos e buscas navegam pelo teclado (setas ↑/↓, Enter para escolher, Esc para fechar).")
B("A associação item→ativo é aprendida como sinônimo: quanto mais você confere, menos correção nas próximas cotações.")
B("O fracionamento mínimo informado pelo fornecedor vira a base de compra; confira sempre a unidade (g × kg).")
B("A conferência do farmacêutico é opcional e não trava a operação — mas é o que garante o comparativo correto.")

doc.save(OUT)
print("Salvo:", OUT)
